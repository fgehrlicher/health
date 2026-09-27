"""Read and validate the pinned BLS 4.0 workbook without touching the database."""

import re
from dataclasses import dataclass
from dataclasses import field as dataclass_field
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK_SHA256 = "524bbefe25b691f5cb3de7a9f3e27fa2967aebfeabf217d99414ba7806e78c60"
SOURCE_NAME = "BLS 4.0"
BLS_GROUP_CODES = frozenset("BCDEFGHKMNPQRSTUVWXY")
# Database column -> BLS component header prefix and unit.
NUTRIENTS = {
    "energy_kj": ("ENERCJ ", "[kJ/100g]"),
    "energy_kcal": ("ENERCC ", "[kcal/100g]"),
    "fat_g": ("FAT ", "[g/100g]"),
    "saturated_fat_g": ("FASAT ", "[g/100g]"),
    "monounsaturated_fat_g": ("FAMS ", "[g/100g]"),
    "polyunsaturated_fat_g": ("FAPU ", "[g/100g]"),
    "carbs_g": ("CHO ", "[g/100g]"),
    "sugars_g": ("SUGAR ", "[g/100g]"),
    "polyols_g": ("POLYL ", "[g/100g]"),
    "starch_g": ("STARCH ", "[g/100g]"),
    "fiber_g": ("FIBT ", "[g/100g]"),
    "protein_g": ("PROT625 ", "[g/100g]"),
    "salt_g": ("NACL ", "[g/100g]"),
    "alcohol_g": ("ALC ", "[g/100g]"),
}
# Read only to correct energy for the BLS 4.0 erratum; not stored.
ERRATUM_INPUTS = {
    "oligosaccharides": ("OLSAC ", "[g/100g]"),
    "organic_acids": ("OA ", "[g/100g]"),
}
# food_source_nutrients key -> BLS header prefix, BLS unit, and the unit stored
# (nutrients.unit). Stored units follow EU labels, so B6 converts from µg to mg.
OTHER_NUTRIENTS = {
    "vitamin_a": ("VITA ", "[µg/100g]", "µg"),
    "vitamin_d": ("VITD ", "[µg/100g]", "µg"),
    "vitamin_e": ("VITE ", "[mg/100g]", "mg"),
    "vitamin_k": ("VITK ", "[µg/100g]", "µg"),
    "vitamin_c": ("VITC ", "[mg/100g]", "mg"),
    "thiamin": ("THIA ", "[mg/100g]", "mg"),
    "riboflavin": ("RIBF ", "[mg/100g]", "mg"),
    "niacin": ("NIAEQ ", "[mg/100g]", "mg"),
    "vitamin_b6": ("VITB6 ", "[µg/100g]", "mg"),
    "folate": ("FOL ", "[µg/100g]", "µg"),
    "vitamin_b12": ("VITB12 ", "[µg/100g]", "µg"),
    "sodium": ("NA ", "[mg/100g]", "mg"),
    "potassium": ("K ", "[mg/100g]", "mg"),
    "calcium": ("CA ", "[mg/100g]", "mg"),
    "magnesium": ("MG ", "[mg/100g]", "mg"),
    "phosphorus": ("P ", "[mg/100g]", "mg"),
    "iron": ("FE ", "[mg/100g]", "mg"),
    "zinc": ("ZN ", "[mg/100g]", "mg"),
    "iodine": ("ID ", "[µg/100g]", "µg"),
    "omega_3": ("FAPUN3 ", "[g/100g]", "g"),
    "epa": ("F20:5CN3 ", "[g/100g]", "g"),
    "dha": ("F22:6CN3 ", "[g/100g]", "g"),
    "omega_6": ("FAPUN6 ", "[g/100g]", "g"),
    "cholesterol": ("CHORL ", "[mg/100g]", "mg"),
    "water": ("WATER ", "[g/100g]", "g"),
    "lactose": ("LACS ", "[g/100g]", "g"),
}
GRAMS_PER_UNIT = {"g": Decimal(1), "mg": Decimal("0.001"), "µg": Decimal("0.000001")}
FIELDS = (
    NUTRIENTS
    | ERRATUM_INPUTS
    | {key: (header, unit) for key, (header, unit, _stored) in OTHER_NUTRIENTS.items()}
)
# BLS component code (INFOODS tagname), e.g. "VITA", -> field above.
COMPONENTS = {header.strip(): field for field, (header, _unit) in FIELDS.items()}
ENERGY_LIMITS = {"energy_kj": 4200, "energy_kcal": 1000}
MARKERS = {"", "TR", "<LOD", "<LOQ", "<LOD or <LOQ", "-"}
# Preparation words in German BLS names -> foods.preparation_state.
PREPARATIONS = {
    "roh": "raw",
    "gekocht": "boiled",
    "gedünstet": "stewed",
    "geschmort": "braised",
    "gegrillt": "grilled",
    "gebraten": "fried",
    "gebacken": "baked",
    "frittiert": "deep-fried",
    "getrocknet": "dried",
    "tiefgefroren": "frozen",
    "Konserve": "canned",
    "geräuchert": "smoked",
    "pochiert": "poached",
    "gargezogen": "poached",
    "gedämpft": "steamed",
    "druckgedämpft": "steamed",
    "geröstet": "roasted",
    "überbacken": "gratinated",
    "getoastet": "toasted",
    "blanchiert": "blanched",
    "gegart": "cooked",
}
PREPARATION_PATTERN = re.compile(
    r"(?<!\w)(" + "|".join(sorted(PREPARATIONS, key=len, reverse=True)) + r")(?!\w)"
)


@dataclass(frozen=True)
class Food:
    code: str
    german_name: str
    english_name: str
    # Keyed by database column, see NUTRIENTS.
    nutrients: dict[str, Decimal | None]
    # Keyed by nutrients.key in stored units, see OTHER_NUTRIENTS.
    other_nutrients: dict[str, Decimal | None]
    preparation_state: str | None


def convert(value: Decimal, from_unit: str, to_unit: str) -> Decimal:
    if from_unit == to_unit:
        return value
    return value * GRAMS_PER_UNIT[from_unit] / GRAMS_PER_UNIT[to_unit]


def preparation_state(german_name: str) -> str | None:
    """The last preparation named outside parentheses: "gekocht, gebraten" is fried.

    Parentheses describe ingredients, e.g. "Karottensalat (gegart) mit Marinade".
    """
    found = PREPARATION_PATTERN.findall(re.sub(r"\([^)]*\)", "", german_name))
    return PREPARATIONS[found[-1]] if found else None


def parse_codes(contents: str) -> set[str]:
    codes = set()
    for line_number, line in enumerate(contents.splitlines(), 1):
        code = line.strip()
        if not code or code.startswith("#"):
            continue
        if len(code) != 7 or not code.isascii() or not code.isalnum() or code != code.upper():
            raise ValueError(f"invalid BLS code {code!r} at line {line_number}")
        if code in codes:
            raise ValueError(f"duplicate BLS code {code} at line {line_number}")
        codes.add(code)
    if not codes:
        raise ValueError("codes file does not select any foods")
    return codes


def columns_from_headers(headers: tuple) -> dict[str, int]:
    def find(name: str, unit: str | None = None) -> int:
        matches = [
            index
            for index, header in enumerate(headers)
            if isinstance(header, str)
            and (header == name if unit is None else header.startswith(name) and unit in header)
        ]
        if len(matches) != 1:
            raise ValueError(
                f"expected one BLS column {name!r} with unit {unit!r}; got {len(matches)}"
            )
        return matches[0]

    columns = {
        "code": find("BLS Code"),
        "german_name": find("Lebensmittelbezeichnung"),
        "english_name": find("Food name"),
    }
    columns.update({field: find(*header) for field, header in FIELDS.items()})
    return columns


def required_text(value: object, field: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field} must be nonempty text")
    return value.strip()


def parse_amount(value: object, code: str, field: str, issues: list[dict]) -> Decimal | None:
    if value is None or isinstance(value, str) and value.strip() in MARKERS:
        marker = "blank" if value is None or not value.strip() else value.strip()
        issues.append({"code": code, "field": field, "marker": marker})
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float, str)):
        raise TypeError(f"unsupported {field} value {value!r}")
    try:
        amount = Decimal(str(value).strip())
    except Exception as error:
        raise ValueError(f"invalid {field} value {value!r}") from error
    if not amount.is_finite() or amount < 0:
        raise ValueError(f"invalid {field} value {value!r}")
    return amount


def corrected_energy(values: dict[str, Decimal]) -> tuple[Decimal, Decimal]:
    """Energy from the BLS 4.0 erratum formula with EU factors, rounded last."""
    available = values["carbs_g"] - values["polyols_g"]
    kcal = (
        values["protein_g"] * 4
        + values["fat_g"] * 9
        + available * 4
        + values["fiber_g"] * 2
        + values["alcohol_g"] * 7
        + values["organic_acids"] * 3
        + values["polyols_g"] * Decimal("2.4")
    )
    kj = (
        values["protein_g"] * 17
        + values["fat_g"] * 37
        + available * 17
        + values["fiber_g"] * 8
        + values["alcohol_g"] * 29
        + values["organic_acids"] * 13
        + values["polyols_g"] * 10
    )
    return tuple(value.quantize(Decimal(1), rounding=ROUND_HALF_UP) for value in (kj, kcal))


def parse_food(
    row: tuple, columns: dict[str, int], code: str, issues: list[dict]
) -> tuple[Food, bool, list[str]]:
    """Parse one row. Impossible optional nutrients become NULL with an "implausible" issue."""
    values = {
        field: parse_amount(row[columns[field]], code, header.strip(), issues)
        for field, (header, _unit) in FIELDS.items()
    }
    german_name = required_text(row[columns["german_name"]], "German food name")
    english_name = required_text(row[columns["english_name"]], "English food name")
    for field, limit in ENERGY_LIMITS.items():
        if values[field] is None:
            raise ValueError(f"{field} must be numeric")
        if values[field] > limit:
            raise ValueError(f"{field} exceeds {limit} per 100 g")
    for field, value in values.items():
        grams = field.endswith("_g") or field in ERRATUM_INPUTS
        if grams and value is not None and value > 100:
            raise ValueError(f"{field} exceeds 100 g per 100 g")
    other = {}
    for key, (header, bls_unit, unit) in OTHER_NUTRIENTS.items():
        value = values[key]
        bls_unit = bls_unit.strip("[]").split("/")[0]
        if value is not None and value * GRAMS_PER_UNIT[bls_unit] > 100:
            # E.g. BLS lists 100.585 g water per 100 g for some diet colas.
            issues.append({"code": code, "field": header.strip(), "marker": "implausible"})
            value = None
        other[key] = None if value is None else convert(value, bls_unit, unit)
    for part in ("sugars_g", "polyols_g"):
        if None not in (values[part], values["carbs_g"]) and values[part] > values["carbs_g"]:
            raise ValueError(f"{part} exceeds carbs_g")

    oligosaccharides = values["oligosaccharides"]
    affected = oligosaccharides is not None and oligosaccharides > 0
    missing_energy_inputs = []
    if affected:
        inputs = {
            "PROT625": "protein_g",
            "FAT": "fat_g",
            "CHO": "carbs_g",
            "FIBT": "fiber_g",
            "ALC": "alcohol_g",
            "OA": "organic_acids",
            "POLYL": "polyols_g",
        }
        missing_energy_inputs = [name for name, field in inputs.items() if values[field] is None]
        if missing_energy_inputs:
            # Published energy is affected by the erratum, but correction inputs
            # are missing. Keep the food and its usable values, not suspect energy.
            values["energy_kj"] = values["energy_kcal"] = None
        else:
            # BLS 4.0 erratum: OLSAC was counted twice in kJ and kcal.
            values["energy_kj"], values["energy_kcal"] = corrected_energy(values)
            for field, limit in ENERGY_LIMITS.items():
                if values[field] < 0 or values[field] > limit:
                    raise ValueError(f"corrected {field} is outside 0–{limit} per 100 g")
    nutrients = {field: values[field] for field in NUTRIENTS}
    return (
        Food(code, german_name, english_name, nutrients, other, preparation_state(german_name)),
        affected and not missing_energy_inputs,
        missing_energy_inputs,
    )


@dataclass
class Reading:
    foods: list[Food]
    # Nonnumeric source markers, e.g. {"code": ..., "field": "FAT", "marker": "TR"}.
    issues: list[dict] = dataclass_field(default_factory=list)
    energy_corrected: int = 0
    energy_unavailable: list[dict] = dataclass_field(default_factory=list)
    # Applied erratum entries: {"code", "components", "action", "section"}.
    erratum_applied: list[dict] = dataclass_field(default_factory=list)


def read_workbook(
    path: Path, selected: set[str] | None = None, erratum: dict | None = None
) -> Reading:
    """Parse the workbook; `erratum` maps BLS codes to corrections (see erratum.py)."""
    from bls4_importer.erratum import apply_erratum

    erratum = erratum or {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        columns = columns_from_headers(next(rows))
        found: dict[str, Food] = {}
        reading = Reading(foods=[])
        for line_number, row in enumerate(rows, 2):
            code = row[columns["code"]]
            if selected is not None and code not in selected:
                continue
            if (
                not isinstance(code, str)
                or len(code) != 7
                or not code.isascii()
                or not code.isalnum()
                or code != code.upper()
                or code[0] not in BLS_GROUP_CODES
            ):
                raise ValueError(f"invalid BLS code {code!r} at worksheet row {line_number}")
            if code in found:
                raise ValueError(f"duplicate selected BLS code {code} in worksheet")
            try:
                food, was_corrected, missing_inputs = parse_food(row, columns, code, reading.issues)
                if code in erratum:
                    food = apply_erratum(food, erratum[code])
            except ValueError as error:
                raise ValueError(f"BLS row {line_number} ({code}): {error}") from error
            found[code] = food
            reading.energy_corrected += was_corrected
            if missing_inputs:
                reading.energy_unavailable.append({"code": code, "missing_inputs": missing_inputs})
            reading.erratum_applied += [
                {
                    "code": code,
                    "components": list(c.components),
                    "action": c.action,
                    "section": c.section,
                }
                for c in erratum.get(code, [])
            ]
        missing = sorted(selected - found.keys()) if selected is not None else []
        if missing:
            raise ValueError(f"selected BLS codes not found in workbook: {missing}")
        in_scope = erratum.keys() if selected is None else erratum.keys() & selected
        if unmatched := sorted(in_scope - found.keys()):
            raise ValueError(f"erratum codes not found in workbook: {unmatched}")
        reading.foods = [found[code] for code in sorted(found)]
        return reading
    finally:
        workbook.close()
